"""Input ingestion and normalization for resumes and job descriptions."""

from dataclasses import dataclass
from pathlib import PurePath
from typing import Final
import io
import re

from app.privacy import PrivacyError, ensure_text_input, sanitize_filename


MAX_INPUT_BYTES: Final = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS: Final = frozenset({".docx", ".pdf", ".txt"})


class IngestionError(ValueError):
    """Raised when an input cannot be safely converted to normalized text."""


@dataclass(frozen=True)
class IngestedDocument:
    """The source text and its normalized representation for later pipeline stages."""

    source_name: str
    source_type: str
    original_text: str
    normalized_text: str


def ingest_text(text: str, *, source_name: str = "pasted-text") -> IngestedDocument:
    """Ingest pasted text while retaining the exact source separately."""
    try:
        ensure_text_input(text, label="Text input")
        safe_source_name = sanitize_filename(source_name)
    except PrivacyError as error:
        raise IngestionError(str(error)) from error

    return _document(safe_source_name, "text", text)


def ingest_file(
    filename: str,
    content: bytes,
    *,
    content_type: str | None = None,
) -> IngestedDocument:
    """Extract supported file content into the common document representation."""
    try:
        filename = sanitize_filename(filename)
    except PrivacyError as error:
        raise IngestionError(str(error)) from error
    if not isinstance(content, bytes):
        raise IngestionError("File content must be bytes.")
    if not content:
        raise IngestionError("File input cannot be empty.")
    if len(content) > MAX_INPUT_BYTES:
        raise IngestionError("File input exceeds the 5 MB limit.")

    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise IngestionError("Unsupported file type. Use PDF, DOCX, or TXT.")

    if extension == ".txt":
        text = _decode_text(content)
    elif extension == ".pdf":
        text = _extract_pdf(content)
    else:
        text = _extract_docx(content)

    if not text.strip():
        raise IngestionError("File does not contain usable text.")

    return _document(filename, extension[1:], text)


def _document(source_name: str, source_type: str, original_text: str) -> IngestedDocument:
    normalized_text = _normalize(original_text)
    if not normalized_text:
        raise IngestionError("Input does not contain usable text.")
    return IngestedDocument(source_name, source_type, original_text, normalized_text)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+$", "", line) for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise IngestionError("TXT files must be UTF-8 encoded.") from error


def _extract_pdf(content: bytes) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as error:
        raise IngestionError("Could not read the PDF file.") from error


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as error:
        raise IngestionError("Could not read the DOCX file.") from error